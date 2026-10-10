"""Fail-closed, source-addressable trade-line extraction for uploaded documents.

This module produces review candidates, never verified customs facts. PDF text
without a stable table row is retained as a document issue rather than joined
to an unrelated number elsewhere on the page.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import time
from decimal import Decimal, InvalidOperation

ALIASES = {
    "shipment_ref": ("shipment_ref", "transaction_id", "line_id", "item_id", "ref"),
    "invoice_number": ("invoice_number", "invoice_no", "invoice", "invoice_num"),
    "cn_code": ("cn_code", "cn", "hs_code", "hsn_code", "commodity_code", "ritc_cth"),
    "quantity": ("quantity", "quantity_t", "weight", "net_weight", "net_mass", "mass"),
    "quantity_unit": ("quantity_unit", "weight_unit", "mass_unit", "unit", "uom"),
    "value": ("value", "line_value", "customs_value", "invoice_value", "amount"),
    "currency": ("currency", "value_currency", "invoice_currency"),
    "origin_country": ("origin_country", "country_of_origin", "origin"),
    "destination_country": ("destination_country", "destination", "import_country"),
    "import_date": ("import_date", "shipment_date", "date"),
    "heat_number": ("heat_number", "heat_no", "heat"),
    "coil_number": ("coil_number", "coil_no", "coil"),
    "container_number": ("container_number", "container_no", "container"),
    "bill_of_lading": ("bill_of_lading", "bl_number", "b_l_number", "bol_number"),
}
COUNTRIES = {"INDIA": "IN", "IN": "IN", "GERMANY": "DE", "DE": "DE",
             "NETHERLANDS": "NL", "NL": "NL", "BELGIUM": "BE", "BE": "BE",
             "FRANCE": "FR", "FR": "FR", "ITALY": "IT", "IT": "IT",
             "UNITED STATES": "US", "USA": "US", "US": "US", "HUNGARY": "HU", "HU": "HU"}
EU = {"AT", "BE", "BG", "HR", "CY", "CZ", "DK", "EE", "FI", "FR", "DE", "GR",
      "HU", "IE", "IT", "LV", "LT", "LU", "MT", "NL", "PL", "PT", "RO", "SK", "SI", "ES", "SE"}


def _key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _number(value: object) -> Decimal | None:
    raw = str(value or "").strip()
    if not re.fullmatch(r"(?:0|[1-9]\d*)(?:\.\d+)?", raw):
        return None
    try:
        return Decimal(raw)
    except InvalidOperation:
        return None


def _country(value: object) -> str | None:
    return COUNTRIES.get(str(value or "").strip().upper())


def _cells_to_line(cells: list[object], headers: list[str], *, filename: str,
                   sha256: str, row_number: int, sheet: str | None = None) -> dict:
    lookup = {_key(h): i for i, h in enumerate(headers) if h is not None}
    found = {}
    provenance = {}
    for field, aliases in ALIASES.items():
        idx = next((lookup[a] for a in aliases if a in lookup), None)
        if idx is None or idx >= len(cells):
            continue
        raw = cells[idx]
        if raw is None or str(raw).strip() == "":
            continue
        found[field] = str(raw).strip()
        provenance[field] = {"document": filename, "sha256": sha256,
                             "sheet": sheet, "row": row_number,
                             "column": idx + 1, "cell": f"{_column(idx + 1)}{row_number}",
                             "raw": str(raw), "page": None, "bbox": None}
    issues = []
    cn = re.sub(r"\s", "", found.get("cn_code", ""))
    if not re.fullmatch(r"\d{8}(?:\d{2})?", cn):
        issues.append("CN_MISSING_OR_INVALID")
    quantity = _number(found.get("quantity"))
    unit = found.get("quantity_unit", "").upper()
    if quantity is None or quantity <= 0:
        issues.append("QUANTITY_MISSING_OR_AMBIGUOUS")
    if unit in {"KG", "KGS"} and quantity is not None:
        quantity_t = quantity / Decimal(1000)
    elif unit in {"T", "MT", "TONNE", "TONNES"}:
        quantity_t = quantity
    else:
        quantity_t = None
        issues.append("WEIGHT_UNIT_MISSING_OR_INVALID")
    value = _number(found.get("value"))
    currency = found.get("currency", "").upper()
    if value is None or value <= 0:
        issues.append("VALUE_MISSING_OR_AMBIGUOUS")
    if currency not in {"EUR", "USD", "INR", "GBP"}:
        issues.append("CURRENCY_MISSING_OR_INVALID")
    if currency != "EUR":
        issues.append("EUR_CONVERSION_REQUIRES_RATE_SNAPSHOT")
    origin = _country(found.get("origin_country"))
    destination = _country(found.get("destination_country"))
    if not origin:
        issues.append("ORIGIN_MISSING_OR_INVALID")
    if not destination:
        issues.append("DESTINATION_MISSING_OR_INVALID")
    elif destination not in EU:
        issues.append("NON_EU_DESTINATION")
    date = found.get("import_date")
    if not date or not re.fullmatch(r"20\d\d-\d\d-\d\d", date):
        issues.append("IMPORT_DATE_MISSING_OR_INVALID")
    if not found.get("invoice_number"):
        issues.append("INVOICE_NUMBER_MISSING")
    return {"shipment_ref": found.get("shipment_ref") or f"{filename}:{sheet or 'sheet'}:{row_number}",
            "invoice_number": found.get("invoice_number"), "cn_code": cn or None,
            "quantity_t": float(quantity_t) if quantity_t is not None else None,
            "customs_value_eur": float(value) if value is not None and currency == "EUR" else None,
            "value": float(value) if value is not None else None, "currency": currency or None,
            "origin_country": origin, "destination_country": destination,
            "import_date": date, "heat_number": found.get("heat_number"),
            "coil_number": found.get("coil_number"),
            "container_number": found.get("container_number"),
            "bill_of_lading": found.get("bill_of_lading"),
            "field_provenance": provenance, "issues": issues,
            "review_status": "REQUIRES_REVIEW"}


def _column(index: int) -> str:
    result = ""
    while index:
        index, rem = divmod(index - 1, 26)
        result = chr(65 + rem) + result
    return result


def ingest_tabular(data: bytes, filename: str) -> dict:
    start = time.perf_counter()
    digest = hashlib.sha256(data).hexdigest()
    rows = []
    issues = []
    if filename.lower().endswith(".csv"):
        try:
            table = list(csv.reader(io.StringIO(data.decode("utf-8-sig"), newline="")))
        except (UnicodeDecodeError, csv.Error) as exc:
            return {"ok": False, "code": "CSV_PARSE_FAILED", "reason": str(exc), "lines": []}
        if table:
            headers = table[0]
            if len({_key(h) for h in headers}) != len(headers):
                issues.append({"code": "DUPLICATE_CSV_HEADERS"})
            for i, row in enumerate(table[1:10001], 2):
                if not any(str(c).strip() for c in row):
                    continue
                line = _cells_to_line(row, headers, filename=filename,
                                      sha256=digest, row_number=i)
                if len(row) != len(headers):
                    line["issues"].append("RAGGED_CSV_ROW")
                rows.append(line)
            if len(table) > 10001:
                issues.append({"code": "ROW_LIMIT_EXCEEDED", "limit": 10000})
    elif filename.lower().endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            for ws in wb.worksheets[:20]:
                iterator = ws.iter_rows(values_only=True)
                headers = next(iterator, None)
                if not headers:
                    continue
                for i, row in enumerate(iterator, 2):
                    if i > 10001:
                        issues.append({"sheet": ws.title, "code": "ROW_LIMIT_EXCEEDED", "limit": 10000})
                        break
                    if not any(c is not None and str(c).strip() for c in row):
                        continue
                    if any(isinstance(c, str) and c.startswith("=") for c in row):
                        issues.append({"sheet": ws.title, "row": i, "code": "FORMULA_REQUIRES_REVIEW"})
                    rows.append(_cells_to_line(list(row), list(headers), filename=filename,
                                               sha256=digest, row_number=i, sheet=ws.title))
            if len(wb.worksheets) > 20:
                issues.append({"code": "SHEET_LIMIT_EXCEEDED", "limit": 20})
            wb.close()
        except Exception as exc:
            return {"ok": False, "code": "XLSX_PARSE_FAILED", "reason": str(exc), "lines": []}
    else:
        return {"ok": False, "code": "UNSUPPORTED_TYPE", "lines": []}
    if not rows:
        issues.append({"code": "NO_DATA_ROWS"})
    for line in rows:
        line["issues"].extend(issue["code"] for issue in issues
                              if issue["code"] not in {"NO_DATA_ROWS"})
    return {"ok": bool(rows), "document_sha256": digest, "lines": rows,
            "issues": issues, "processing_time_ms": round((time.perf_counter() - start) * 1000, 1)}


def ingest_pdf(data: bytes, filename: str) -> dict:
    """Preserve page evidence; PDF text is review-only until line joins exist."""
    start = time.perf_counter()
    digest = hashlib.sha256(data).hexdigest()
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = [{"page": i, "text": p.extract_text() or ""}
                 for i, p in enumerate(reader.pages, 1)]
    except Exception as exc:
        return {"ok": False, "code": "PDF_PARSE_FAILED", "reason": str(exc), "lines": []}
    issues = []
    if sum(len(p["text"].strip()) for p in pages) < 20:
        issues.append({"code": "PDF_OCR_REQUIRED"})
    # Only an explicitly labelled, self-contained line is offered for review.
    # No document-wide first-CN / largest-weight / largest-value joins.
    pattern = re.compile(
        r"(?i)\b(?:CN|HS|HSN)\s*(?:CODE)?\s*[:#-]?\s*(?P<cn>\d{8}(?:\d{2})?)"
        r"\s+QTY\s*[:#-]?\s*(?P<qty>\d+(?:\.\d+)?)\s*(?P<unit>KG|KGS|MT|T|TONNES?)"
        r"\s+(?:VALUE|AMOUNT)\s*[:#-]?\s*(?P<currency>EUR|USD|INR|GBP)\s*(?P<value>\d+(?:\.\d+)?)"
    )
    lines = []
    for page in pages:
        for line_no, raw_line in enumerate(page["text"].splitlines(), 1):
            match = pattern.search(raw_line)
            if not match:
                continue
            row = [None, match["cn"], match["qty"], match["unit"], match["value"],
                   match["currency"], None, None, None]
            headers = ["invoice_number", "cn_code", "quantity", "quantity_unit",
                       "value", "currency", "origin_country", "destination_country", "import_date"]
            candidate = _cells_to_line(row, headers, filename=filename, sha256=digest,
                                       row_number=line_no)
            for field in ("cn_code", "quantity", "quantity_unit", "value", "currency"):
                candidate["field_provenance"][field].update(
                    {"page": page["page"], "line": line_no, "cell": None, "bbox": None})
            candidate["issues"].append("PDF_BOUNDING_BOX_UNAVAILABLE")
            lines.append(candidate)
    if not lines:
        issues.append({"code": "PDF_LINE_ASSOCIATION_UNVERIFIED",
                       "detail": "No explicit CN/QTY/VALUE line; table/row join requires review."})
    return {"ok": bool(lines), "document_sha256": digest, "pages": len(pages),
            "lines": lines, "issues": issues,
            "processing_time_ms": round((time.perf_counter() - start) * 1000, 1)}
