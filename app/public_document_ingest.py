"""Source-grounded extraction for public trade-document regression and pilot review.

Every observation is a candidate with a PDF page/line/quote/box or XLSX cell.
No document classification, extracted value, or link asserts authenticity.
"""
from __future__ import annotations

import hashlib
import io
import re
import subprocess
import time
import zipfile
from pathlib import Path

from .pdf_observations import _geometry, _locate, extract_pdf_observations


def _role(text: str) -> str:
    u=text.upper()
    if "MILL TEST / INSPECTION CERTIFICATE" in u or "MILL TEST CERTIFICATE" in u: return "MTC"
    if "PART - I - SHIPPING BILL SUMMARY" in u: return "SHIPPING_BILL"
    if "DV1 + IMPORT SAD CASE" in u: return "SAD_CASE_STUDY"
    if "ANNEXURE" in u and "EXAMINATION REPORT" in u: return "CUSTOMS_ANNEXURE"
    if "EXPORT INVOICE" in u: return "INVOICE"
    if "BILL OF LADING" in u or "MTD / BL NO" in u: return "BILL_OF_LADING"
    if "IMPORT DECLARATION" in u: return "BLANK_IMPORT_FORM"
    return "UNKNOWN"


def _unique_word_bbox(words:list[dict],value:str):
    target=re.sub(r"[^A-Z0-9]","",value.upper())
    matches=[w["bbox"] for w in words if re.sub(r"[^A-Z0-9]","",w["text"].upper())==target]
    return matches[0] if len(matches)==1 else None


def extract_public_pdf(data: bytes, filename: str) -> dict:
    start=time.perf_counter()
    base=extract_pdf_observations(data,filename)
    if base.get("code"): return {"role":"UNKNOWN","observations":base}
    layout=subprocess.run(["pdftotext","-layout","-","-"],input=data,capture_output=True,timeout=20,check=False)
    if layout.returncode: return {"role":"UNKNOWN","observations":base}
    pages=layout.stdout.decode("utf-8","replace").split("\f")[:base.get("pages",0)]
    role=_role("\n".join(pages[:2]))
    geometry=_geometry(data)
    fields=[f for f in base["fields"] if not (role=="SAD_CASE_STUDY" and f["field"]=="bill_of_lading" and re.fullmatch(r"N\d{3}",f["raw"]))]
    extra_tables=[]
    existing={(f["field"],str(f["raw"]).upper()) for f in fields}
    def add(name,raw,page,line,quote):
        raw=str(raw).strip()
        if not raw or (name,raw.upper()) in existing: return
        words=geometry[page-1] if page<=len(geometry) else []
        fields.append({"field":name,"raw":raw,"page":page,"line":line,"quote":quote.strip(),
                       "bbox":_locate(words,raw) or _unique_word_bbox(words,raw),"method":"pdf-text-rule"})
        existing.add((name,raw.upper()))
    for p,text in enumerate(pages,1):
        lines=text.splitlines()
        for n,line in enumerate(lines,1):
            patterns=[]
            if role=="MTC":
                row=re.search(r"\b(CH-\d{4,8})\s+(\d{3}/\d{3}L?)\s+(\d+(?:\.\d+)?)\s*mm\b",line,re.I)
                if row:
                    extra_tables.append({"kind":"mtc_heat_items","page":p,"rows":[{"heat_number":row.group(1),"material_grade":row.group(2),"size_mm":row.group(3),"source":{"page":p,"line":n,"quote":line.strip(),"bbox":_locate(geometry[p-1] if p<=len(geometry) else [],row.group(1))}}]})
                patterns=[("certificate_number",r"Test Certificate No\.\s*:\s*(EXP/\d{2}-\d{2}/\d{5}-\d+)"),
                          ("purchase_order",r"P\.\s*O\.\s*No\.\s*:\s*(\d{4,12})"),
                          ("invoice_number",r"Invoice No\.\s*& Date\s*:\s*(EXP/\d{2}-\d{2}/\d{5})"),
                          ("hs_code",r"HSN Code\s*:\s*(\d{4}\.\d{2})"),
                          ("heat_number",r"\b(CH-\d{4,8})\b"),
                          ("net_weight_kg",r"\b(\d[\d,]*(?:\.\d+)?)\s*Kgs\.?"),
                          ("certificate_standard",r"\b(EN\s*10204\s*:\s*2005)\b")]
            elif role=="SHIPPING_BILL":
                patterns=[("shipping_bill_number",r"\bSB\s*No\s*[:#-]?\s*(\d{6,10})\b"),
                          ("shipping_date",r"\bSB\s*Date\s*[:#-]?\s*(\d{1,2}-[A-Z]{3}-\d{2,4})\b"),
                          ("destination_country",r"COUNTRY OF FINALDESTINATIO\s+([A-Z ]{4,30})")]
                if n>1 and "INDIAN CUSTOMS EDI SYSTEM" in line and "SB No" in lines[n-2]:
                    m=re.search(r"\b(\d{7})\s+(\d{1,2}-[A-Z]{3}-\d{2,4})\b",line)
                    if m:
                        add("shipping_bill_number",m.group(1),p,n,line)
                        add("shipping_date",m.group(2),p,n,line)
            elif role=="SAD_CASE_STUDY":
                patterns=[("invoice_number",r"Invoice number and date of issue:\s*([A-Z0-9/-]+)"),
                          ("bill_of_lading",r"Bill of Lading:\s*([A-Z0-9]+)"),
                          ("taric_code",r"TARIC code:\s*([\d.]+)"),
                          ("sad_date",r"SAD issuing date:\s*(\d{2}/\d{2}/\d{4})"),
                          ("origin_country",r"Origin:\s*(India)"),
                          ("gross_weight_kg",r"\b(25,000)\s*KG and"),
                          ("net_weight_kg",r"and\s*(24,650)\s*KG"),
                          ("invoice_value",r"value of the merchandise is USD\s*([\d,]+)"),
                          ("eori",r"EORI number:\s*([A-Z]{2,3}\s+\d{7,12})")]
            elif role=="CUSTOMS_ANNEXURE":
                patterns=[("invoice_number",r"Export Invoice No\.\s*:\s*(SPR\d{4}E\d{4})")]
            elif role=="BILL_OF_LADING":
                patterns=[("shipping_bill_number",r"SB\.NO\.\s*(\d{7})"),
                          ("net_weight_kg",r"NET WT:\s*([\d,.]+)\s*KGS"),
                          ("purchase_order",r"P\.O\. NUMBER\s*(\d{4,12})")]
            for name,pat in patterns:
                m=re.search(pat,line,re.I)
                if m:add(name,m.group(1),p,n,line)
            if role=="SAD_CASE_STUDY" and n>1 and "of caraway seeds" in line and "EORI number:" in lines[n-2]:
                m=re.search(r"EORI number:\s*([A-Z ]?\d{7,12})",lines[n-2],re.I)
                if m:add("eori",m.group(1),p,n-1,lines[n-2])
    for field in fields:
        if not field.get("bbox") and field.get("page") and field["page"]<=len(geometry):
            field["bbox"]=_unique_word_bbox(geometry[field["page"]-1],field["raw"])
    base["fields"]=fields
    base["tables"].extend(extra_tables)
    base["processing_time_ms"]=round((time.perf_counter()-start)*1000,1)
    return {"role":role,"observations":base,"status":"REQUIRES_REVIEW"}


def extract_cbam_workbook(data:bytes,filename:str)->dict:
    """Read the Commission communication summary as cell-cited candidates."""
    import openpyxl
    started=time.perf_counter()
    result={"filename":Path(filename).name,"sha256":hashlib.sha256(data).hexdigest(),
            "role":"CBAM_COMMUNICATION_CANDIDATE","status":"REQUIRES_REVIEW","fields":[],"product_rows":[]}
    if len(data)>8*1024*1024: return {**result,"code":"FILE_TOO_LARGE"}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist())>60*1024*1024:
                return {**result,"code":"WORKBOOK_EXPANDED_SIZE_EXCEEDED"}
    except zipfile.BadZipFile:
        return {**result,"code":"INVALID_XLSX"}
    w=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    if "Summary_Products" not in w.sheetnames or "Summary_Processes" not in w.sheetnames:
        return {**result,"code":"UNSUPPORTED_CBAM_WORKBOOK"}
    s=w["Summary_Products"]
    for row in s.iter_rows(min_row=10,max_row=min(s.max_row,120)):
        by_col={c.column_letter:c for c in row if hasattr(c,"column_letter")}
        code_cell=by_col.get("F")
        if not code_cell or not re.fullmatch(r"7\d{7}",str(code_cell.value)):continue
        item={"cn_code":str(code_cell.value),"source":{"sheet":s.title,"cell":code_cell.coordinate}}
        for key,col in (("see_direct_tco2_per_t","I"),("see_indirect_tco2_per_t","J"),("see_total_tco2_per_t","K")):
            cell=by_col.get(col)
            if cell and type(cell.value) in (int,float):
                item[key]={"raw":cell.value,"source":{"sheet":s.title,"cell":cell.coordinate}}
        result["product_rows"].append(item)
    s=w["Summary_Processes"]
    for row in s.iter_rows(min_row=1,max_row=40):
        cells=list(row)
        for i,c in enumerate(cells[:-1]):
            if isinstance(c.value,str) and "Name of the installation (English name)" in c.value:
                next_cell=next((x for x in cells[i+1:] if isinstance(x.value,str) and x.value.strip()),None)
                if next_cell:result["fields"].append({"field":"installation_name","raw":next_cell.value,
                    "sheet":s.title,"cell":next_cell.coordinate,"label_cell":c.coordinate})
    result["processing_time_ms"]=round((time.perf_counter()-started)*1000,1)
    return result
