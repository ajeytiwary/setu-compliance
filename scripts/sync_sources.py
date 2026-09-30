"""Unified sync entry point: auto resolvers + weekly/manual versioned pipeline.

Two complementary pipelines (both fail-closed, engines consume normalized only):

1. AUTO (remote, CI-safe) -- `app/source_sync.py::sync_many`
   Positional datasets, link-discovery resolvers for the 5 pluggable sources:
   taric_measures, eucdm, echa_candidate_list, scip_schema, eu_sanctions.
   Used by `.github/workflows/sync-regulatory-sources.yml` (`--strict`).

2. WEEKLY/MANUAL (versioned) -- `app/source_resolvers.py::publish_normalized`
   `--dataset/--all/--file/--url/--as-of/--force/--list` for all 11 registry
   datasets incl. CBAM xlsx, steel 2026/1457, quota, Comext, demo telemetry.
   Bot-blocked publishers (ECHA 403, FSF auth) ingest via `--file` browser
   download. Layout: data/raw/<ds>/<YYYY-MM-DD-sha8>/ + data/normalized/... +
   data/manifests/<ds>-latest.json. `--all` never stops at first failure,
   writes `_sync_report.json`, SKIPPED_FRESH idempotency for weekly cron.

Usage:
  python scripts/sync_sources.py --strict                    # CI (5 auto sources)
  python scripts/sync_sources.py --list                      # download guidance
  python scripts/sync_sources.py --dataset eu_sanctions --as-of 2026-09-29
  python scripts/sync_sources.py --dataset echa_candidate_list --file ~/dl/candidate_list_en.csv
  python scripts/sync_sources.py --all                       # weekly cron entry
  python scripts/sync_sources.py --all --force               # re-fetch today
"""
from __future__ import annotations
import argparse
import io
import json
import sys
import urllib.request
import zipfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app import source_resolvers as R  # noqa: E402
from app.data_sources import selected  # noqa: E402
from app.source_sync import sync_many as auto_sync_many  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

AUTO_DEFAULT = ["taric_measures", "eucdm", "echa_candidate_list", "scip_schema", "eu_sanctions"]

# Direct versioned distributions. `url=None` means the publisher only offers
# a landing page / auth-gated API: download in a browser, then --file it.
SOURCES: dict[str, dict] = {
    "taric_measures": {
        "provider_id": "eu_taxud_taric_bulk",
        "url": None,
        "where": "Daily TARIC delta ZIP from the taric-opendata mirror "
                 "(https://github.com/rousseauxy/taric-opendata, newest eu-* TARIC_<date>.zip), "
                 "or export CSV/XML from https://taxation-customs.ec.europa.eu/online-services/online-services-and-databases-customs/eu-customs-tariff-taric_en, "
                 "then --file it. AUTO path: set SETU_TARIC_BULK_URL or use positional sync.",
        "filename": "taric-delta.zip", "content_type": "application/zip",
        "authority": "MIRROR", "legal_authority": False, "min_records": 1,
    },
    "quota_balances": {
        "provider_id": "eu_taxud_quota_export",
        "url": None,
        "where": "Export CSV from the QUOTA consultation database "
                 "(https://taxation-customs.ec.europa.eu/customs/common-customs-tariff-cct/tariff-quotas_en), then --file it. Refresh daily.",
        "filename": "quota.csv", "content_type": "text/csv",
        "authority": "OFFICIAL", "legal_authority": False, "min_records": 1,
    },
    "eucdm": {
        "provider_id": "eu_taxud_eucdm_annex_b",
        "url": "https://eucdm.softdev.eu.com/EUCDM/Download/EUCDM-HTML_v7p0p11_2026-08-19.zip",
        "where": "EUCDM v7.0.11 HTML distribution (softdev mirror of DG TAXUD; needs a browser "
                 "User-Agent). Annex-B declaration data elements parsed from EN/EUCDM/Annex-B/sd1.htm. "
                 "AUTO path: positional `python scripts/sync_sources.py eucdm` fetches it directly.",
        "filename": "eucdm-7.0.11.zip",
        "content_type": "application/zip",
        "authority": "MIRROR", "legal_authority": False, "min_records": 1,
    },
    "steel_2026_1457": {
        "provider_id": "eurlex_2026_1457",
        "url": "https://eur-lex.europa.eu/legal-content/EN/TXT/?uri=CELEX:32026R1457",
        "where": "Implementing Regulation (EU) 2026/1457, ELI https://eur-lex.europa.eu/eli/reg_impl/2026/1457/oj/eng",
        "filename": "steel-2026-1457.html", "content_type": "text/html",
        "authority": "LEGAL", "legal_authority": True, "min_records": 1,
    },
    "cbam_defaults": {
        "provider_id": "eu_taxud_cbam_defaults",
        "url": None,  # resolved from registry at runtime (TAXUD document UUIDs rotate)
        "where": "Registry provider eu_taxud_cbam_defaults (TAXUD CBAM defaults workbook).",
        "filename": "cbam-defaults.xlsx",
        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "authority": "OFFICIAL", "legal_authority": False, "min_records": 1,
    },
    "cbam_benchmarks": {
        "provider_id": "eu_taxud_cbam_benchmarks",
        "url": None,
        "where": "Registry provider eu_taxud_cbam_benchmarks (TAXUD benchmarks workbook).",
        "filename": "cbam-benchmarks.xlsx",
        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "authority": "OFFICIAL", "legal_authority": False, "min_records": 1,
    },
    "echa_candidate_list": {
        "provider_id": "echa_candidate_csv",
        "url": "https://raw.githubusercontent.com/analeonescu/chemical-security-evals/main/data/chemicals_databases/candidate-list-of-svhc-for-authorisation-export.csv",
        "where": "Tab-delimited Candidate List mirror (ECHA blocks bots with 403). "
                 "AUTO path: positional sync fetches the mirror directly (~507 substances).",
        "filename": "candidate_list_en.csv", "content_type": "text/csv",
        "authority": "MIRROR", "legal_authority": False, "min_records": 1,
    },
    "scip_schema": {
        "provider_id": "echa_scip_610",
        "url": "https://raw.githubusercontent.com/USEPA/CompTox-IUCLIDTools/dev/entity_models/article_6_8/models/article_9_0.py",
        "where": "IUCLID ARTICLE.9.0 dossier models (SCIP 6.10 schema/picklists) from "
                 "USEPA/CompTox-IUCLIDTools dev branch -- the official ECHA SCIP 6.10 ZIP "
                 "(https://echa.europa.eu/en/scip-format) blocks automated fetches (Azure WAF), "
                 "so the versioned path fetches the three model files directly; "
                 "a browser-downloaded official ZIP can still be ingested via --file.",
        "filename": "article_9_0.py", "content_type": "text/x-python",
        "authority": "MIRROR", "legal_authority": False, "min_records": 1,
    },
    "eu_sanctions": {
        "provider_id": "eu_fsf_11_xml",
        "url": "https://webgate.ec.europa.eu/fsd/fsf/public/files/xmlFullSanctionsList_1_1/content?token=dG9rZW4tMjAxNw",
        "where": "EU FSF 1.1 consolidated XML resolved via the data.europa.eu DCAT record "
                 "(distributions carry download_url/access_url token links). "
                 "Versions are date-foldered and never overwritten. "
                 "AUTO path: positional sync uses DCAT discovery with webgate fallback.",
        "filename": "eu-sanctions-map-xml-all.xml", "content_type": "application/xml",
        "authority": "OFFICIAL", "legal_authority": False, "min_records": 1,
    },
    "comext_trade": {
        "provider_id": "eurostat_comext",
        "url": "https://ec.europa.eu/eurostat/api/comext/dissemination",
        "where": "Eurostat Comext dissemination API; filtered India->EU chapters 72/73, 2024-2026.",
        "filename": "comext.json", "content_type": "application/json",
        "authority": "OFFICIAL_STATISTICS", "legal_authority": False, "min_records": 0,
    },
    "industrial_telemetry_demo": {
        "provider_id": "uci_steel_energy",
        "url": None,
        "where": "UCI Steel Industry Energy Consumption (DEMO_ONLY) https://archive.ics.uci.edu/dataset/851/steel+industry+energy+consumption -- download CSV, then --file it.",
        "filename": "steel-energy.csv", "content_type": "text/csv",
        "authority": "DEMO_ONLY", "legal_authority": False, "min_records": 1,
    },
}


def _xlsx_rows(raw: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    out, seen = [], set()
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(x).strip() if x is not None else "" for x in rows[0]]
        for row in rows[1:]:
            rec = {header[i] or f"col_{i}": v for i, v in enumerate(row)
                   if v is not None and str(v).strip() != ""}
            if not rec:
                continue
            key = json.dumps({"sheet": ws.title, **rec}, sort_keys=True, default=str)
            if key in seen:
                continue
            seen.add(key)
            out.append({"sheet": ws.title, **rec})
    return out


def _resolve_url(dataset: str) -> str | None:
    spec = SOURCES[dataset]
    if spec["url"]:
        return spec["url"]
    if dataset in ("cbam_defaults", "cbam_benchmarks"):
        try:
            return selected(dataset, spec["provider_id"]).url
        except Exception:
            return None
    return None


def sync_dataset(dataset: str, as_of: str | None = None, url: str | None = None,
                 file: str | None = None, force: bool = False) -> dict:
    """Versioned weekly/manual sync for one dataset (raw -> normalized -> manifest)."""
    if dataset not in SOURCES:
        raise ValueError(f"Unknown dataset: {dataset}")
    spec = SOURCES[dataset]
    as_of = as_of or date.today().isoformat()

    if not force and not file:
        cur = R.latest_manifest(dataset)
        if cur and cur.get("effective_from") == as_of:
            return {"dataset": dataset, "status": "SKIPPED_FRESH", "manifest": cur}

    if file:
        raw = Path(file).read_bytes()
        src = f"file://{Path(file).resolve()}"
        ctype = spec["content_type"]
    else:
        target = url or _resolve_url(dataset)
        if not target:
            raise ValueError(f"{dataset}: no direct download URL. {spec['where']}")
        try:
            # EUCDM's mirror sits behind CloudFront: needs a browser User-Agent.
            if dataset == "eucdm":
                import urllib.request as _url
                req = _url.Request(target, headers={"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36"})
                with _url.urlopen(req, timeout=120) as _r:
                    raw = _r.read()
                    ctype = (_r.headers.get("Content-Type") or spec["content_type"]).split(";")[0].strip()
            else:
                raw, ctype = R.fetch_url(target)
        except Exception as e:
            raise ValueError(f"{dataset}: download failed ({e}). {spec['where']} "
                             f"Workaround: download in a browser, then "
                             f"`python scripts/sync_sources.py --dataset {dataset} --file <path>`")
        src = target

    if dataset == "taric_measures":
        if raw[:2] == b"PK":
            from app.source_sync import normalize_taric as _nt
            records = _nt(raw)  # daily delta ZIP (Measures/MEAS_TYP_ID dialect)
        elif raw.lstrip()[:1] == b"<":
            records = R.normalize_taric_xml(raw)
            if not records:
                raise ValueError("taric_measures: no measure rows parsed; supply a TARIC CSV/XML export via --file")
        else:
            records = R.normalize_taric_csv(raw.decode("utf-8-sig"), src, as_of)["records"]
    elif dataset == "quota_balances":
        from app.eu_public_data import normalize_quota
        import csv as _csv
        records = normalize_quota(list(_csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))), as_of)["records"]
    elif dataset == "eucdm":
        if raw[:2] == b"PK":
            from app.source_sync import normalize_eucdm as _ne
            records = _ne(raw)  # EUCDM HTML distribution ZIP (sd1.htm D.E.s)
        else:
            records = R.normalize_eucdm_rows(_xlsx_rows(raw))
    elif dataset == "steel_2026_1457":
        records = R.normalize_steel_1457(raw.decode("utf-8", "ignore"), src)
        cats = records[0]["categories"] if records else []
        if len(cats) < 26:
            raise RuntimeError(f"steel annex incomplete ({len(cats)} categories); checked-in legal snapshot stays authoritative")
    elif dataset in ("cbam_defaults", "cbam_benchmarks"):
        records = _xlsx_rows(raw)
    elif dataset == "echa_candidate_list":
        records = R.normalize_echa_candidate_csv(raw.decode("utf-8-sig", "ignore"))
    elif dataset == "scip_schema":
        if raw[:2] == b"PK":
            inv = R.normalize_scip_zip(raw)
            records = [{"version": "6.10", "file": f} for f in inv["files"]] + \
                      [{"version": "6.10", "picklist": p["list"], "value": p["value"]}
                       for p in inv["picklist_values"][:5000]]
        else:
            # Versioned path: fetch the sibling model files next to the
            # resolved article_9_0.py (or parse a --file'd model file alone).
            import urllib.request as _url
            text = raw.decode("utf-8", "ignore")
            files = {"article_9_0.py": text}
            if not file:
                base = (url or _resolve_url(dataset) or "").rsplit("/", 1)[0]
                for sib in ("common_types_domain_v9.py", "platform_fields.py"):
                    try:
                        req = _url.Request(base + "/" + sib,
                                           headers={"User-Agent": "SetuCompliance/0.7 (+weekly regulatory sync)"})
                        with _url.urlopen(req, timeout=90) as _r:
                            files[sib] = _r.read().decode("utf-8", "ignore")
                    except Exception:
                        pass
            records = R.normalize_iuclid_article_models(files)
    elif dataset == "eu_sanctions":
        if raw.lstrip()[:1] == b"<":
            records = R.normalize_sanctions_xml(raw)
        else:
            records = R.normalize_sanctions_csv(raw.decode("utf-8-sig", "ignore"))
    elif dataset == "comext_trade":
        obj = json.loads(raw.decode("utf-8", "ignore"))
        records = R.normalize_comext(obj)
    elif dataset == "industrial_telemetry_demo":
        import csv as _csv
        records = list(_csv.DictReader(io.StringIO(raw.decode("utf-8-sig", "ignore"))))
    else:
        raise ValueError(dataset)

    manifest = R.publish_normalized(
        dataset, spec["provider_id"], src, raw, records,
        content_type=ctype, authority=spec["authority"],
        legal_authority=spec["legal_authority"], as_of=as_of,
        min_records=spec["min_records"], filename=spec["filename"])
    bridge_to_engines(dataset, manifest)
    return {"dataset": dataset, "status": "SYNCED", "manifest": manifest}


def bridge_to_engines(dataset: str, manifest: dict) -> None:
    """Push normalized snapshots into the live engine caches (never raw)."""
    npath = R.NORMALIZED / dataset / manifest["version"] / "normalized.json"
    norm = json.loads(npath.read_text())
    recs = norm["records"]
    if dataset == "taric_measures":
        from app.eu_public_data import normalize_taric, store_snapshot
        store_snapshot("taric", normalize_taric(recs, norm["as_of"]),
                       json.dumps(recs, sort_keys=True).encode())
    elif dataset == "quota_balances":
        from app.eu_public_data import normalize_quota, store_snapshot
        store_snapshot("quota", normalize_quota(recs, norm["as_of"]),
                       json.dumps(recs, sort_keys=True).encode())


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Download, hash, normalize and validate Setu regulatory sources")
    ap.add_argument("datasets", nargs="*", default=None,
                    help="AUTO positional datasets (default: 5 pluggable sources via link-discovery)")
    ap.add_argument("--strict", action="store_true", help="exit non-zero if any source fails (CI)")
    ap.add_argument("--dataset", help="Versioned single-dataset sync (weekly/manual pipeline)")
    ap.add_argument("--all", action="store_true", help="Versioned sync of all datasets (weekly cron)")
    ap.add_argument("--list", action="store_true", help="List datasets + download guidance")
    ap.add_argument("--file", help="Manual file to ingest (browser download)")
    ap.add_argument("--url", help="Override distribution URL")
    ap.add_argument("--as-of", dest="as_of", help="Effective date YYYY-MM-DD (default today)")
    ap.add_argument("--force", action="store_true", help="Re-fetch even if today's manifest exists")
    args = ap.parse_args(argv)

    if args.list:
        for k, s in SOURCES.items():
            print(f"{k}\n  provider: {s['provider_id']}\n  url: {_resolve_url(k)}\n  how: {s['where']}\n")
        return 0

    # Versioned weekly/manual path takes precedence when its flags are used.
    if args.all or args.dataset or args.file or args.url or args.as_of or args.force:
        targets = list(SOURCES) if args.all else ([args.dataset] if args.dataset else [])
        if not targets:
            ap.error("pass --dataset ID or --all (see --list)")
        if args.file and len(targets) > 1:
            ap.error("--file only works with a single --dataset")
        report = {"as_of": args.as_of or date.today().isoformat(), "results": {}}
        code = 0
        for ds in targets:
            try:
                r = sync_dataset(ds, as_of=args.as_of, url=args.url,
                                 file=args.file, force=args.force)
                report["results"][ds] = {"status": r["status"],
                                         "sha256": r["manifest"]["sha256"][:12],
                                         "records": r["manifest"]["record_count"]}
                print(f"OK {ds}: {r['status']} records={r['manifest']['record_count']} sha={r['manifest']['sha256'][:12]}")
            except Exception as e:
                report["results"][ds] = {"status": "FAILED", "error": str(e)[:300]}
                print(f"FAIL {ds}: {e}")
                code = 1
                if not args.all:
                    break
        R.MANIFESTS.mkdir(parents=True, exist_ok=True)
        (R.MANIFESTS / "_sync_report.json").write_text(json.dumps(report, indent=2))
        return code

    # AUTO path (remote behaviour preserved): positional link-discovery resolvers.
    datasets = args.datasets or AUTO_DEFAULT
    r = auto_sync_many(datasets)
    print(json.dumps(r, indent=2, default=str))
    if args.strict and not all(x["ok"] for x in r.values()):
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
