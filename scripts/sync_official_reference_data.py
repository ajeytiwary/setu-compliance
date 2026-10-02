from __future__ import annotations
import hashlib,io,json,re,urllib.request,zipfile
from pathlib import Path
from openpyxl import load_workbook
from app.reference_validation import validate
ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/"data"/"official"; STAGE=OUT/"staging"; OUT.mkdir(parents=True,exist_ok=True);STAGE.mkdir(parents=True,exist_ok=True)
from app.data_sources import selected
DEFAULT_URL=selected("cbam_defaults").url
BENCHMARK_URL=selected("cbam_benchmarks").url
STEEL_URL=selected("steel_2026_1457").url
EXAMPLES_URL=selected("cbam_examples").url
def fetch(url):
 req=urllib.request.Request(url,headers={"User-Agent":"EuroSetuCompliance/0.5"});return urllib.request.urlopen(req,timeout=60).read()
def xlsx_rows(raw):
 wb=load_workbook(io.BytesIO(raw),data_only=True,read_only=True);out=[];seen=set()
 for ws in wb.worksheets:
  rows=list(ws.iter_rows(values_only=True))
  if not rows:continue
  header=[str(x).strip() if x is not None else "" for x in rows[0]]
  for row in rows[1:]:
   rec={header[i] or "col_"+str(i):v for i,v in enumerate(row) if v is not None and str(v).strip()!=""}
   if not rec:continue
   key=json.dumps({"sheet":ws.title,**rec},sort_keys=True,default=str)
   if key in seen:continue
   seen.add(key); out.append({"sheet":ws.title,**rec})
 return out
def publish(name,source,raw,records,min_records=1):
 check=validate(name,raw,records,min_records)
 obj={"source":source,"sha256":check["sha256"],"validation":check,"records":records}
 (STAGE/name).write_text(json.dumps(obj,indent=2,default=str))
 if not check["valid"]:raise RuntimeError(name+" failed validation: "+",".join(check["errors"]))
 (OUT/name).write_text(json.dumps(obj,indent=2,default=str));return check
def sync_cbam():
 return {name:publish(name,url,(raw:=fetch(url)),xlsx_rows(raw),1) for name,url in (("cbam-defaults-official.json",DEFAULT_URL),("cbam-benchmarks-official.json",BENCHMARK_URL))}
def sync_examples():
 raw=fetch(EXAMPLES_URL); z=zipfile.ZipFile(io.BytesIO(raw)); records=[]
 for name in sorted(z.namelist()):
  if not name.lower().endswith(".xlsx"):continue
  b=z.read(name)
  try:
   wb=load_workbook(io.BytesIO(b),data_only=False,read_only=True)
   records.append({"filename":name,"sha256":hashlib.sha256(b).hexdigest(),
                   "sheets":wb.sheetnames,"sheet_count":len(wb.sheetnames)})
  except Exception as e:
   raise RuntimeError("official CBAM example workbook unreadable: "+name+": "+str(e))
 if len(records)<7:
  raise RuntimeError("official CBAM example bundle incomplete: expected >=7 xlsx workbooks, got "+str(len(records)))
 return publish("cbam-communication-examples-official.json",EXAMPLES_URL,raw,records,7)

def sync_steel():
 raw=fetch(STEEL_URL);html=raw.decode("utf-8","ignore")
 cats=sorted(set(re.findall(r'>\s*(\d{1,2}(?:\.[AB])?)\s*<',html)));orders=sorted(set(re.findall(r'09\.\d{4}',html)));cns=sorted(set("".join(x.split()) for x in re.findall(r'\b(?:72|73)\d{2}(?:\s*\d{2}){2}\b',html)))
 records=[{"categories":cats,"order_numbers":orders,"cn_codes":cns}]
 check=publish("steel-2026-1457-extract.json",STEEL_URL,raw,records,1)
 if len(cats)<26:raise RuntimeError("steel annex incomplete: expected at least 26 category labels")
 return check
if __name__=="__main__":print(json.dumps({"cbam":sync_cbam(),"cbam_examples":sync_examples(),"steel":sync_steel()},indent=2))
